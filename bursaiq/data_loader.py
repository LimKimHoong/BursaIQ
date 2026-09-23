"""Validated local Excel/PDF/JSON ingestion for the BursaIQ Stage 02 demo."""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


TOKEN_PATTERN = re.compile(r"[a-z0-9][a-z0-9-]{1,}", re.IGNORECASE)
QUERY_STOPWORDS = {
    "about", "are", "bursa", "can", "could", "current", "does", "explain", "for", "from", "give", "how", "in", "is",
    "language", "me", "of", "plain", "please", "show", "summarise", "summarize", "tell", "the", "this",
    "to", "what", "when", "where", "which", "who", "why", "with",
}
ROLE_WORKSPACES = {
    "gcmc": {"market", "learn"},
    "securities": {"market", "learn", "reg"},
    "hr": {"market", "hr", "learn"},
    "finance": {"market", "learn"},
}


class IngestionError(RuntimeError):
    """Raised when a controlled input cannot be validated."""


@dataclass(frozen=True)
class SearchResult:
    document_id: str
    title: str
    filename: str
    workspace: str
    excerpt: str
    score: int

    def as_dict(self) -> dict[str, Any]:
        return {
            "documentId": self.document_id,
            "title": self.title,
            "filename": self.filename,
            "workspace": self.workspace,
            "excerpt": self.excerpt,
            "score": self.score,
        }


class LocalDataRepository:
    """Load the manifest and its approved local source files into a small in-memory index."""

    def __init__(self, input_dir: Path) -> None:
        self.input_dir = input_dir.resolve()
        self.manifest_path = self.input_dir / "manifest.json"
        self._fingerprint: tuple[tuple[str, int, int], ...] = ()
        self._cache: dict[str, Any] | None = None

    def _safe_path(self, relative: str) -> Path:
        candidate = (self.input_dir / relative).resolve()
        if self.input_dir not in candidate.parents:
            raise IngestionError(f"Source path escapes Input directory: {relative}")
        if candidate.suffix.lower() not in {".xlsx", ".pdf", ".csv", ".json"}:
            raise IngestionError(f"Unsupported source type: {candidate.suffix}")
        return candidate

    def _current_fingerprint(self) -> tuple[tuple[str, int, int], ...]:
        files = [self.manifest_path]
        if self.manifest_path.exists():
            try:
                manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
                files.extend(self._safe_path(item["filename"]) for item in manifest.get("documents", []))
            except (json.JSONDecodeError, KeyError, TypeError, IngestionError):
                pass
        return tuple(
            sorted(
                (str(path), path.stat().st_mtime_ns, path.stat().st_size)
                for path in files
                if path.exists()
            )
        )

    def load(self, force: bool = False) -> dict[str, Any]:
        fingerprint = self._current_fingerprint()
        if not force and self._cache is not None and fingerprint == self._fingerprint:
            return self._cache
        if not self.manifest_path.exists():
            raise IngestionError("Input/manifest.json is missing. Run scripts/seed_demo_data.py.")
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        if manifest.get("classification") != "SYNTHETIC_DEMO_ONLY":
            raise IngestionError("Manifest classification must be SYNTHETIC_DEMO_ONLY for Stage 02.")

        documents: list[dict[str, Any]] = []
        market: dict[str, Any] = {}
        hr: dict[str, Any] = {"procedure": [], "applications": []}
        errors: list[dict[str, str]] = []

        for item in manifest.get("documents", []):
            try:
                document = self._load_document(item)
                documents.append(document)
                if item.get("id") == "gcmc-pulse":
                    market = self._load_market_workbook(self._safe_path(item["filename"]))
                elif item.get("id") == "hr-applications":
                    hr["applications"] = self._read_table(self._safe_path(item["filename"]), "Applications")
                elif item.get("id") == "hr-procedure":
                    hr["procedure"] = self._parse_hr_procedure(document.get("_text", ""))
            except Exception as exc:  # keep the demo available while surfacing a specific source failure
                errors.append({"source": str(item.get("filename", "unknown")), "error": str(exc)})

        if not market:
            raise IngestionError("The GCMC market workbook could not be loaded.")
        if not hr["procedure"]:
            hr["procedure"] = self._default_hr_procedure()

        self._cache = {
            "meta": {
                "asOf": "31 Jul 2026",
                "datasetLabel": "Synthetic competition dataset",
                "classification": manifest["classification"],
                "loadedAt": datetime.now(timezone.utc).isoformat(),
                "sourceCount": len(documents),
                "ingestionErrors": errors,
            },
            "market": market,
            "hr": hr,
            "documents": documents,
        }
        self._fingerprint = fingerprint
        return self._cache

    def public_bootstrap(self) -> dict[str, Any]:
        loaded = self.load()
        return {
            "meta": loaded["meta"],
            "market": loaded["market"],
            "hr": loaded["hr"],
            "documents": [{key: value for key, value in doc.items() if not key.startswith("_")} for doc in loaded["documents"]],
        }

    def search(self, query: str, workspace: str, role: str, limit: int = 5) -> list[SearchResult]:
        if workspace not in ROLE_WORKSPACES.get(role, set()):
            return []
        all_tokens = set(TOKEN_PATTERN.findall(query.lower()))
        tokens = all_tokens.difference(QUERY_STOPWORDS) or all_tokens
        if not tokens:
            return []
        results: list[SearchResult] = []
        for doc in self.load()["documents"]:
            if doc["workspace"] != workspace:
                continue
            haystack = f'{doc["title"]} {doc.get("excerpt", "")} {doc.get("_text", "")}'.lower()
            haystack_tokens = set(TOKEN_PATTERN.findall(haystack))
            matches = [
                token for token in tokens
                if token in haystack_tokens
                or (not token.endswith("s") and f"{token}s" in haystack_tokens)
                or (token.endswith("s") and token[:-1] in haystack_tokens)
            ]
            if not matches:
                continue
            excerpt = self._best_excerpt(doc.get("_text", doc.get("excerpt", "")), tokens)
            results.append(SearchResult(doc["id"], doc["title"], doc["filename"], doc["workspace"], excerpt, len(matches)))
        return sorted(results, key=lambda item: (-item.score, item.title))[: max(1, min(limit, 10))]

    def source_path(self, document_id: str, role: str) -> Path:
        for document in self.load()["documents"]:
            if document["id"] != document_id:
                continue
            if document["workspace"] not in ROLE_WORKSPACES.get(role, set()):
                raise PermissionError(document_id)
            return self._safe_path(document["_relativePath"])
        raise KeyError(document_id)

    def _load_document(self, item: dict[str, Any]) -> dict[str, Any]:
        path = self._safe_path(item["filename"])
        if not path.exists():
            raise IngestionError(f"Approved source is missing: {item['filename']}")
        document = dict(item)
        document["_relativePath"] = item["filename"]
        document["filename"] = item.get("displayFilename", path.name)
        document["bytes"] = path.stat().st_size
        if path.suffix.lower() == ".pdf":
            reader = PdfReader(str(path))
            document["pages"] = f"{len(reader.pages)} pages"
            document["_text"] = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif path.suffix.lower() == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=True)
            document["pages"] = f"{len(book.sheetnames)} worksheet{'s' if len(book.sheetnames) != 1 else ''}"
            text_parts: list[str] = []
            for sheet in book.worksheets:
                text_parts.append(sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    text_parts.append(" | ".join("" if value is None else str(value) for value in row))
            document["_text"] = "\n".join(text_parts)
            book.close()
        elif path.suffix.lower() == ".json":
            payload = json.loads(path.read_text(encoding="utf-8"))
            topics = payload.get("topics", []) if isinstance(payload, dict) else []
            text_parts = [str(payload.get("title", document["title"]))] if isinstance(payload, dict) else []
            for topic in topics:
                if not isinstance(topic, dict):
                    continue
                text_parts.extend([str(topic.get("title", "")), str(topic.get("summary", ""))])
                text_parts.extend(str(action) for action in topic.get("actions", []))
            if isinstance(payload, dict):
                text_parts.append(str(payload.get("disclaimer", "")))
            document["pages"] = f"{len(topics)} demo topics"
            document["_text"] = "\n".join(part for part in text_parts if part)
        return document

    @staticmethod
    def _read_table(path: Path, sheet_name: str) -> list[dict[str, Any]]:
        book = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in book.sheetnames:
            book.close()
            raise IngestionError(f"Worksheet '{sheet_name}' not found in {path.name}")
        sheet = book[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) for value in next(rows)]
        records = [{headers[index]: value for index, value in enumerate(row)} for row in rows if any(value is not None for value in row)]
        book.close()
        return records

    def _load_market_workbook(self, path: Path) -> dict[str, Any]:
        book = load_workbook(path, read_only=True, data_only=True)
        required = {"Headline", "Monthly", "Sectors", "Counters", "Participation", "Regional"}
        missing = required.difference(book.sheetnames)
        book.close()
        if missing:
            raise IngestionError(f"Market workbook is missing worksheets: {', '.join(sorted(missing))}")
        headline_rows = self._read_table(path, "Headline")
        headline = {str(row["Metric"]): row["Value"] for row in headline_rows}
        integer_metrics = {"gainers", "losers", "unchanged"}
        for key in integer_metrics:
            if key in headline:
                headline[key] = int(headline[key])
        return {
            "headline": headline,
            "monthly": self._read_table(path, "Monthly"),
            "sectors": self._read_table(path, "Sectors"),
            "counters": self._read_table(path, "Counters"),
            "participation": self._read_table(path, "Participation"),
            "regional": self._read_table(path, "Regional"),
        }

    @staticmethod
    def _parse_hr_procedure(text: str) -> list[dict[str, Any]]:
        stages: list[dict[str, Any]] = []
        blocks = re.split(r"(?=\d+\.\s+(?:Requisition|Sourcing|Panel|Pre-employment|Offer))", text)
        for block in blocks:
            heading = re.search(r"^(\d+)\.\s+([^\n]+)", block.strip())
            owner = re.search(r"Owner:\s*([^\.]+)", block)
            target = re.search(r"Target:\s*(\d+)\s+working", block)
            if heading and owner and target:
                stages.append({"step": int(heading.group(1)), "name": heading.group(2).strip(), "owner": owner.group(1).strip(), "targetDays": int(target.group(1))})
        return stages

    @staticmethod
    def _default_hr_procedure() -> list[dict[str, Any]]:
        return [
            {"step": 1, "name": "Requisition approval", "owner": "Hiring manager & Finance", "targetDays": 2},
            {"step": 2, "name": "Sourcing and screening", "owner": "Talent Acquisition", "targetDays": 8},
            {"step": 3, "name": "Panel assessment", "owner": "Hiring panel", "targetDays": 5},
            {"step": 4, "name": "Pre-employment checks", "owner": "Talent Acquisition", "targetDays": 4},
            {"step": 5, "name": "Offer approval and issue", "owner": "HR approver", "targetDays": 3},
        ]

    @staticmethod
    def _best_excerpt(text: str, tokens: set[str], length: int = 900) -> str:
        collapsed = " ".join(text.split())
        lowered = collapsed.lower()
        positions = []
        for token in tokens:
            variants = {token, token[:-1] if token.endswith("s") else f"{token}s"}
            matches = [re.search(rf"\b{re.escape(variant)}\b", lowered) for variant in variants if variant]
            positions.extend(match.start() for match in matches if match)
        start = max(0, (min(positions) if positions else 0) - 70)
        excerpt = collapsed[start : start + length].strip()
        if start:
            excerpt = "…" + excerpt
        if start + length < len(collapsed):
            excerpt += "…"
        return excerpt
