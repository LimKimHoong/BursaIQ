"""Permission-aware local document indexing and retrieval for Ask Bursa."""

from __future__ import annotations

import csv
import hashlib
import html
import json
import math
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".csv", ".txt", ".md", ".html", ".htm"}
DEPARTMENTS = {"human_resources", "market_operations", "regulations", "group_technology", "corporate_policies"}
DEFAULT_ROLES = ["gcmc", "hr", "securities", "finance"]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class KnowledgeIndex:
    def __init__(self, database_path: Path, knowledge_root: Path, embed: Callable[[list[str]], list[list[float]]], suggest: Callable[[str, dict[str, Any]], list[dict[str, str]]] | None = None) -> None:
        self.database_path = database_path
        self.knowledge_root = knowledge_root.resolve()
        self.embed = embed
        self.suggest = suggest
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_documents (
                    id TEXT PRIMARY KEY, relative_path TEXT UNIQUE NOT NULL, filename TEXT NOT NULL,
                    department TEXT NOT NULL, classification TEXT NOT NULL, allowed_roles_json TEXT NOT NULL,
                    owner TEXT NOT NULL, effective_date TEXT, version TEXT NOT NULL, sha256 TEXT NOT NULL,
                    indexed_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id TEXT PRIMARY KEY, document_id TEXT NOT NULL, ordinal INTEGER NOT NULL,
                    location TEXT NOT NULL, text_content TEXT NOT NULL, embedding_json TEXT NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES knowledge_documents(id) ON DELETE CASCADE
                );
                CREATE INDEX IF NOT EXISTS knowledge_chunks_document ON knowledge_chunks(document_id, ordinal);
                CREATE TABLE IF NOT EXISTS knowledge_suggestions (
                    id TEXT PRIMARY KEY, document_id TEXT NOT NULL, question TEXT NOT NULL,
                    question_key TEXT NOT NULL, description TEXT NOT NULL, department TEXT NOT NULL,
                    allowed_roles_json TEXT NOT NULL, source_version TEXT NOT NULL, generated_at TEXT NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES knowledge_documents(id) ON DELETE CASCADE
                );
                CREATE INDEX IF NOT EXISTS knowledge_suggestions_role ON knowledge_suggestions(department, generated_at DESC);
                CREATE TABLE IF NOT EXISTS knowledge_audit (
                    id TEXT PRIMARY KEY, occurred_at TEXT NOT NULL, role TEXT NOT NULL, question TEXT NOT NULL,
                    document_ids_json TEXT NOT NULL, outcome TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS knowledge_ingestion_runs (
                    id TEXT PRIMARY KEY, started_at TEXT NOT NULL, completed_at TEXT NOT NULL,
                    indexed_count INTEGER NOT NULL, unchanged_count INTEGER NOT NULL,
                    removed_count INTEGER NOT NULL, errors_json TEXT NOT NULL
                );
                """
            )

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def ingest(self) -> dict[str, Any]:
        started = utc_now()
        indexed = unchanged = removed = 0
        errors: list[dict[str, str]] = []
        seen: set[str] = set()
        for path in self._source_files():
            relative = path.relative_to(self.knowledge_root).as_posix()
            seen.add(relative)
            try:
                metadata_path = path.with_name(path.name + ".metadata.json")
                digest_input = path.read_bytes() + (metadata_path.read_bytes() if metadata_path.exists() else b"")
                digest = hashlib.sha256(digest_input).hexdigest()
                existing = self._document_by_path(relative)
                if existing and existing["sha256"] == digest and (not self.suggest or self._suggestion_count(existing["id"]) > 0):
                    unchanged += 1
                    continue
                sections = extract_sections(path)
                chunks = chunk_sections(sections)
                if not chunks:
                    raise ValueError("No readable text was extracted.")
                embeddings = self.embed([chunk["text"] for chunk in chunks])
                if len(embeddings) != len(chunks):
                    raise ValueError("Embedding count did not match the extracted chunks.")
                metadata = self._metadata(path)
                suggestions: list[dict[str, str]] = []
                if self.suggest:
                    try:
                        suggestions = validate_suggestions(self.suggest("\n\n".join(chunk["text"] for chunk in chunks)[:12_000], metadata))
                    except Exception as exc:
                        errors.append({"path": relative, "error": f"Question suggestion generation failed: {exc}"})
                self._replace_document(relative, path.name, digest, metadata, chunks, embeddings, suggestions)
                indexed += 1
            except Exception as exc:
                errors.append({"path": relative, "error": str(exc)})
        with self._connect() as connection:
            rows = connection.execute("SELECT relative_path FROM knowledge_documents").fetchall()
            for row in rows:
                if row["relative_path"] not in seen:
                    connection.execute("DELETE FROM knowledge_documents WHERE relative_path = ?", (row["relative_path"],))
                    removed += 1
            run_id = f"ING-{uuid.uuid4().hex[:10].upper()}"
            completed = utc_now()
            connection.execute("INSERT INTO knowledge_ingestion_runs VALUES (?, ?, ?, ?, ?, ?, ?)", (run_id, started, completed, indexed, unchanged, removed, json.dumps(errors)))
        return {"id": run_id, "startedAt": started, "completedAt": completed, "indexed": indexed, "unchanged": unchanged, "removed": removed, "errors": errors}

    def search(self, question: str, role: str, limit: int = 6) -> list[dict[str, Any]]:
        query_embeddings = self.embed([question])
        if not query_embeddings:
            return []
        query = query_embeddings[0]
        with self._connect() as connection:
            rows = connection.execute(
                """SELECT c.*, d.filename, d.department, d.classification, d.allowed_roles_json,
                          d.owner, d.version, d.relative_path
                   FROM knowledge_chunks c JOIN knowledge_documents d ON d.id = c.document_id"""
            ).fetchall()
        ranked = []
        for row in rows:
            if role not in json.loads(row["allowed_roles_json"]):
                continue
            score = cosine(query, json.loads(row["embedding_json"]))
            ranked.append({
                "documentId": row["document_id"], "chunkId": row["id"], "title": row["filename"],
                "department": display_department(row["department"]), "classification": row["classification"],
                "owner": row["owner"], "version": row["version"], "location": row["location"],
                "text": row["text_content"], "score": round(score, 6), "relativePath": row["relative_path"],
            })
        ranked.sort(key=lambda item: item["score"], reverse=True)
        return [item for item in ranked if item["score"] >= 0.15][:limit]

    def record_query(self, role: str, question: str, results: list[dict[str, Any]], outcome: str) -> None:
        document_ids = list(dict.fromkeys(item["documentId"] for item in results))
        with self._connect() as connection:
            connection.execute("INSERT INTO knowledge_audit VALUES (?, ?, ?, ?, ?, ?)", (f"ASK-{uuid.uuid4().hex[:10].upper()}", utc_now(), role, question, json.dumps(document_ids), outcome))

    def status(self) -> dict[str, Any]:
        with self._connect() as connection:
            documents = connection.execute("SELECT COUNT(*) FROM knowledge_documents").fetchone()[0]
            chunks = connection.execute("SELECT COUNT(*) FROM knowledge_chunks").fetchone()[0]
            suggestions = connection.execute("SELECT COUNT(*) FROM knowledge_suggestions").fetchone()[0]
            run = connection.execute("SELECT * FROM knowledge_ingestion_runs ORDER BY completed_at DESC LIMIT 1").fetchone()
        return {"documents": documents, "chunks": chunks, "suggestions": suggestions, "supportedExtensions": sorted(SUPPORTED_EXTENSIONS), "departments": sorted(DEPARTMENTS), "lastIngestion": dict(run) if run else None}

    def list_suggestions(self, role: str, limit: int = 8) -> list[dict[str, Any]]:
        with self._connect() as connection:
            rows = connection.execute(
                """SELECT s.*, d.filename FROM knowledge_suggestions s
                   JOIN knowledge_documents d ON d.id = s.document_id
                   ORDER BY s.generated_at DESC, s.rowid DESC"""
            ).fetchall()
        grouped: dict[str, dict[str, Any]] = {}
        for row in rows:
            if role not in json.loads(row["allowed_roles_json"]):
                continue
            item = grouped.get(row["question_key"])
            if item:
                if row["document_id"] not in item["supportingDocumentIds"]:
                    item["supportingDocumentIds"].append(row["document_id"])
                continue
            grouped[row["question_key"]] = {
                "id": row["id"], "question": row["question"], "description": row["description"],
                "department": display_department(row["department"]), "sourceVersion": row["source_version"],
                "supportingDocumentIds": [row["document_id"]], "sourceDocument": row["filename"],
                "generatedAt": row["generated_at"],
            }
        by_department: dict[str, list[dict[str, Any]]] = {}
        for item in grouped.values():
            by_department.setdefault(item["department"], []).append(item)
        balanced: list[dict[str, Any]] = []
        while by_department and len(balanced) < limit:
            for department in list(by_department):
                items = by_department[department]
                if items:
                    balanced.append(items.pop(0))
                if not items:
                    del by_department[department]
                if len(balanced) >= limit:
                    break
        return balanced

    def document_path(self, document_id: str, role: str) -> Path:
        with self._connect() as connection:
            row = connection.execute("SELECT relative_path, allowed_roles_json FROM knowledge_documents WHERE id = ?", (document_id,)).fetchone()
        if row is None or role not in json.loads(row["allowed_roles_json"]):
            raise KeyError(document_id)
        path = (self.knowledge_root / row["relative_path"]).resolve()
        if self.knowledge_root not in path.parents or not path.is_file():
            raise KeyError(document_id)
        return path

    def _source_files(self) -> list[Path]:
        if not self.knowledge_root.exists():
            return []
        return sorted(path for path in self.knowledge_root.rglob("*") if path.is_file() and not path.name.startswith((".", "_")) and path.suffix.lower() in SUPPORTED_EXTENSIONS)

    def _metadata(self, path: Path) -> dict[str, Any]:
        department = path.relative_to(self.knowledge_root).parts[0] if path.relative_to(self.knowledge_root).parts else "corporate_policies"
        if department not in DEPARTMENTS:
            raise ValueError(f"Document must be inside one of the configured department folders: {', '.join(sorted(DEPARTMENTS))}.")
        metadata_path = path.with_name(path.name + ".metadata.json")
        values: dict[str, Any] = {}
        if metadata_path.exists():
            values = json.loads(metadata_path.read_text(encoding="utf-8"))
            if not isinstance(values, dict):
                raise ValueError("Metadata sidecar must contain a JSON object.")
        roles = values.get("allowed_roles", DEFAULT_ROLES)
        if not isinstance(roles, list) or not roles or not all(isinstance(item, str) and item.strip() for item in roles):
            raise ValueError("allowed_roles must be a non-empty string array.")
        normalized_roles = [item.strip().lower() for item in roles]
        if "employee" in normalized_roles:
            normalized_roles.extend(DEFAULT_ROLES)
        return {"department": department, "classification": str(values.get("classification", "Internal")), "allowed_roles": sorted(set(normalized_roles)), "owner": str(values.get("document_owner", display_department(department))), "effective_date": str(values.get("effective_date", "")) or None, "version": str(values.get("version", "1.0"))}

    def _document_by_path(self, relative: str) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute("SELECT * FROM knowledge_documents WHERE relative_path = ?", (relative,)).fetchone()

    def _suggestion_count(self, document_id: str) -> int:
        with self._connect() as connection:
            return int(connection.execute("SELECT COUNT(*) FROM knowledge_suggestions WHERE document_id = ?", (document_id,)).fetchone()[0])

    def _replace_document(self, relative: str, filename: str, digest: str, metadata: dict[str, Any], chunks: list[dict[str, str]], embeddings: list[list[float]], suggestions: list[dict[str, str]]) -> None:
        existing = self._document_by_path(relative)
        document_id = existing["id"] if existing else f"DOC-{uuid.uuid4().hex[:12].upper()}"
        with self._connect() as connection:
            connection.execute("DELETE FROM knowledge_chunks WHERE document_id = ?", (document_id,))
            connection.execute("DELETE FROM knowledge_suggestions WHERE document_id = ?", (document_id,))
            connection.execute(
                """INSERT INTO knowledge_documents VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                   ON CONFLICT(id) DO UPDATE SET filename=excluded.filename, department=excluded.department,
                   classification=excluded.classification, allowed_roles_json=excluded.allowed_roles_json,
                   owner=excluded.owner, effective_date=excluded.effective_date, version=excluded.version,
                   sha256=excluded.sha256, indexed_at=excluded.indexed_at""",
                (document_id, relative, filename, metadata["department"], metadata["classification"], json.dumps(metadata["allowed_roles"]), metadata["owner"], metadata["effective_date"], metadata["version"], digest, utc_now()),
            )
            for ordinal, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
                connection.execute("INSERT INTO knowledge_chunks VALUES (?, ?, ?, ?, ?, ?)", (f"CHK-{uuid.uuid4().hex[:12].upper()}", document_id, ordinal, chunk["location"], chunk["text"], json.dumps(embedding)))
            for suggestion in suggestions:
                question_key = normalize_question(suggestion["question"])
                connection.execute("INSERT INTO knowledge_suggestions VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", (f"SUG-{uuid.uuid4().hex[:12].upper()}", document_id, suggestion["question"], question_key, suggestion["description"], metadata["department"], json.dumps(metadata["allowed_roles"]), metadata["version"], utc_now()))


def extract_sections(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return [(f"Page {index}", page.extract_text() or "") for index, page in enumerate(PdfReader(str(path)).pages, start=1)]
    if suffix == ".docx":
        try:
            from docx import Document
        except ModuleNotFoundError as exc:
            raise ValueError("Word extraction requires python-docx. Run pip install -r requirements.txt.") from exc
        document = Document(str(path))
        return [("Document", "\n".join(paragraph.text for paragraph in document.paragraphs))]
    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ModuleNotFoundError as exc:
            raise ValueError("PowerPoint extraction requires python-pptx. Run pip install -r requirements.txt.") from exc
        presentation = Presentation(str(path))
        return [(f"Slide {index}", "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))) for index, slide in enumerate(presentation.slides, start=1)]
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        return [(f"Sheet {sheet.title}", "\n".join(" | ".join(str(value) for value in row if value is not None) for row in sheet.iter_rows(values_only=True))) for sheet in workbook.worksheets]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return [("CSV", "\n".join(" | ".join(row) for row in csv.reader(handle)))]
    text = path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        text = html.unescape(re.sub(r"<[^>]+>", " ", text))
    return [("Document", text)]


def chunk_sections(sections: list[tuple[str, str]], size: int = 1200, overlap: int = 180) -> list[dict[str, str]]:
    chunks = []
    for location, raw in sections:
        text = re.sub(r"\s+", " ", raw).strip()
        start = 0
        while text and start < len(text):
            end = min(len(text), start + size)
            if end < len(text):
                boundary = text.rfind(" ", start + size // 2, end)
                if boundary > start:
                    end = boundary
            value = text[start:end].strip()
            if value:
                chunks.append({"location": location, "text": value})
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks


def cosine(left: list[float], right: list[float]) -> float:
    if len(left) != len(right) or not left:
        return 0.0
    denominator = math.sqrt(sum(value * value for value in left)) * math.sqrt(sum(value * value for value in right))
    return sum(a * b for a, b in zip(left, right)) / denominator if denominator else 0.0


def display_department(value: str) -> str:
    return value.replace("_", " ").title().replace("Human Resources", "Human Resources").replace("Group Technology", "Group Technology")


def normalize_question(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def validate_suggestions(values: Any) -> list[dict[str, str]]:
    if not isinstance(values, list):
        raise ValueError("Suggestion generator must return an array.")
    result = []
    seen = set()
    for value in values[:8]:
        if not isinstance(value, dict):
            continue
        question = str(value.get("question") or "").strip()
        description = str(value.get("description") or "").strip()
        key = normalize_question(question)
        if len(question) < 8 or len(question) > 180 or not question.endswith("?") or not key or key in seen:
            continue
        seen.add(key)
        result.append({"question": question, "description": description[:220] or "Suggested from an indexed internal document."})
    return result
