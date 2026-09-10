"""Create the controlled synthetic Excel and PDF source pack for BursaIQ."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "Input"
NAVY = "07171D"
TEAL = "39D3BC"
GOLD = "E6BB63"
PALE = "E9F4F1"


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    thin = Side(style="thin", color="D9E5E2")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def add_sheet(book: Workbook, name: str, headers: list[str], rows: list[list[object]], widths: list[int]) -> None:
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, widths)


def create_market_workbook() -> None:
    path = INPUT / "GCMC" / "GCMC_Market_Pulse.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "Headline", ["Metric", "Value", "Unit", "As of", "Classification"], [
        ["fbmKLCI", 1638.2, "points", "31 Jul 2026", "SYNTHETIC"],
        ["klciMtdPct", 2.4, "%", "31 Jul 2026", "SYNTHETIC"],
        ["klciYtdPct", 5.8, "%", "31 Jul 2026", "SYNTHETIC"],
        ["marketCapBn", 2148.6, "RM bn", "31 Jul 2026", "SYNTHETIC"],
        ["marketCapMtdPct", 2.1, "%", "31 Jul 2026", "SYNTHETIC"],
        ["adv30dBn", 3.42, "RM bn", "31 Jul 2026", "SYNTHETIC"],
        ["advPrior30dBn", 3.08, "RM bn", "17 Jun 2026", "SYNTHETIC"],
        ["velocityPct", 40.1, "%", "31 Jul 2026", "SYNTHETIC"],
        ["velocityPriorPct", 36.1, "%", "30 Jun 2026", "SYNTHETIC"],
        ["gainers", 612, "counters", "31 Jul 2026", "SYNTHETIC"],
        ["losers", 438, "counters", "31 Jul 2026", "SYNTHETIC"],
        ["unchanged", 486, "counters", "31 Jul 2026", "SYNTHETIC"]
    ], [24, 15, 15, 18, 18])
    add_sheet(book, "Monthly", ["month", "index", "valueBn", "volumeBn"], [
        ["Jan", 1557.4, 2.82, 3.31], ["Feb", 1571.8, 2.91, 3.42], ["Mar", 1596.3, 3.05, 3.62],
        ["Apr", 1582.7, 2.94, 3.44], ["May", 1604.9, 3.13, 3.71], ["Jun", 1599.8, 3.08, 3.65],
        ["Jul", 1638.2, 3.42, 4.08]
    ], [15, 16, 16, 16])
    add_sheet(book, "Sectors", ["name", "mtdPct", "contributionPoints", "valueBn"], [
        ["Technology", 6.8, 9.6, 0.54], ["Financial Services", 3.1, 8.2, 0.79], ["Utilities", 4.4, 5.1, 0.23],
        ["Plantation", 1.7, 2.4, 0.19], ["Healthcare", -1.9, -1.8, 0.25]
    ], [24, 14, 22, 14])
    add_sheet(book, "Counters", ["name", "ticker", "contributionPoints", "pricePct"], [
        ["Satria Bank", "SATRIA", 5.4, 4.7], ["Maju Utilities", "MAJU", 3.9, 6.2],
        ["Nusa Digital", "NUSA", 3.1, 8.8], ["Sentral Holdings", "SNTRL", 2.2, 3.4]
    ], [24, 14, 22, 14])
    add_sheet(book, "Participation", ["group", "netFlowMn", "sharePct"], [
        ["Local institutions", 486, 39.4], ["Foreign investors", 218, 24.8], ["Local retail", -704, 35.8]
    ], [24, 18, 16])
    add_sheet(book, "Regional", ["market", "mtdPct", "ytdPct", "currency"], [
        ["Malaysia · FBM KLCI", 2.4, 5.8, "MYR"], ["Singapore · STI", 1.6, 7.1, "SGD"],
        ["Indonesia · JCI", -0.8, 3.5, "IDR"], ["Thailand · SET", 0.9, -2.2, "THB"],
        ["United States · S&P 500", 1.9, 9.6, "USD"]
    ], [32, 14, 14, 14])
    book.properties.title = "BursaIQ Synthetic GCMC Market Pulse"
    book.properties.description = "Invented figures for a competition demonstration. Not official market data."
    book.save(path)


def create_hr_workbook() -> None:
    path = INPUT / "HR" / "HR_Applications_Demo.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Applications"
    sheet.append(["applicant", "ref", "position", "stage", "nextAction", "owner", "classification"])
    rows = [
        ["Alya Rahman", "DEM-26031", "Market Insights Analyst", "Panel assessment", "Second interview · 5 Aug 2026", "N. Hassan", "FICTIONAL"],
        ["Daniel Lim", "DEM-26032", "Market Insights Analyst", "Sourcing and screening", "Hiring manager review · 2 Aug 2026", "N. Hassan", "FICTIONAL"],
        ["Siti Hajar", "DEM-26018", "People Analytics Executive", "Pre-employment checks", "Reference check in progress", "F. Lee", "FICTIONAL"],
        ["Kavin Raj", "DEM-26011", "Cybersecurity Specialist", "Offer approval and issue", "Approval due · 1 Aug 2026", "M. Wong", "FICTIONAL"]
    ]
    for row in rows:
        sheet.append(row)
    style_sheet(sheet, [21, 15, 28, 25, 34, 16, 18])
    book.properties.title = "BursaIQ Fictional Applicant Tracker"
    book.properties.description = "Fictional records for a competition demonstration."
    book.save(path)


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=28, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=12, alignment=TA_LEFT),
        "h1": ParagraphStyle("H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=20, textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=8, spaceAfter=8),
        "body": ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=14, textColor=colors.HexColor("#30464B"), spaceAfter=7),
        "small": ParagraphStyle("Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=11, textColor=colors.HexColor("#60787C"), spaceAfter=5),
        "callout": ParagraphStyle("Callout", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=9, leading=13, textColor=colors.HexColor("#12695E"), backColor=colors.HexColor(f"#{PALE}"), borderPadding=8, spaceBefore=7, spaceAfter=10)
    }


def doc_page(canvas, document) -> None:
    canvas.saveState()
    canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
    canvas.rect(0, A4[1] - 15 * mm, A4[0], 15 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor(f"#{TEAL}"))
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(18 * mm, A4[1] - 10 * mm, "BURSAIQ  /  CONTROLLED DEMO SOURCE")
    canvas.setFillColor(colors.HexColor("#6B7E82"))
    canvas.setFont("Helvetica", 7)
    canvas.drawString(18 * mm, 10 * mm, "CONTROLLED COMPETITION MATERIAL - NOT AN OFFICIAL BURSA MALAYSIA DOCUMENT")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"PAGE {document.page}")
    canvas.restoreState()


def create_pdf(path: Path, title: str, subtitle: str, sections: list[tuple[str, list[str]]], purpose: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()
    purpose_text = purpose or "This controlled source was created for the BursaIQ Stage 02 competition prototype. All examples are invented. Replace it only with material approved for the demonstration environment."
    story = [Spacer(1, 10 * mm), Paragraph(title, styles["title"]), Paragraph(subtitle, styles["callout"]), Paragraph("Purpose and status", styles["h1"]), Paragraph(purpose_text, styles["body"])]
    for index, (heading, paragraphs) in enumerate(sections):
        if index and index % 3 == 0:
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["h1"]))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, styles["body"]))
    story.append(Spacer(1, 7 * mm))
    story.append(Paragraph("Document owner: Demo Content Owner  ·  Review cycle: Before each showcase  ·  Classification: SYNTHETIC_DEMO_ONLY", styles["small"]))
    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=21 * mm, bottomMargin=17 * mm, title=title, author="BursaIQ Demo Team")
    document.build(story, onFirstPage=doc_page, onLaterPages=doc_page)


def create_product_overview_pdf() -> None:
    create_pdf(
        INPUT / "Learn" / "Bursa_Products_Overview.pdf",
        "Bursa Products Overview",
        "High-level product map for new joiners",
        [
            ("Product map at a glance", [
                "Bursa's product choices span securities such as shares, structured products, ETFs, REITs, bonds and sukuk; commodity, equity and financial futures and options; Shariah-compliant participation through Bursa Malaysia-i and Bursa Suq Al-Sila'; indices; the Labuan International Financial Exchange (LFX); and Bursa Gold Dinar."
            ]),
            ("1. Securities market", [
                "Common product groups include shares, structured products such as warrants, exchange-traded funds (ETFs), real estate investment trusts (REITs), and exchange-traded bonds and sukuk.",
                "These products have different structures, risks and eligibility requirements. This overview names categories; it is not investment advice or a complete product specification."
            ]),
            ("2. Derivatives market", [
                "Bursa Malaysia Derivatives provides commodity, equity and financial derivatives. Examples include crude palm oil futures, equity-index futures, single-stock futures, and Malaysian Government Securities futures.",
                "If a user asks specifically about options, examples include equity-index options and options linked to crude palm oil futures. Contract availability and specifications must be checked against the current official product page."
            ]),
            ("3. Islamic market", [
                "Bursa Malaysia-i supports end-to-end Shariah-compliant securities investing. Bursa Suq Al-Sila' is a commodity trading platform that supports Islamic liquidity management and financing."
            ]),
            ("4. Other product and market areas", [
                "The broader product map also includes Bursa Malaysia indices, the Labuan International Financial Exchange (LFX), and Bursa Gold Dinar. Services such as listing, trading, clearing, settlement, depository and market information support these markets."
            ]),
            ("5. How to use this answer", [
                "Start with the user's objective: investing, hedging, Shariah-compliant participation, issuer access, or learning. Then open the relevant official product page for current eligibility, risks, fees and contract specifications.",
                "Public reference: Bursa Malaysia, Our Products and Services - www.bursamalaysia.com/trade/our_products_services. Public reference: Bursa Assist, Bursa Securities Market - assist.bursamalaysia.com."
            ])
        ],
        purpose="This controlled source is a concise public-product summary prepared for the BursaIQ Stage 02 competition prototype. It is not an exhaustive catalogue and must not be treated as investment advice. Product availability and specifications should be checked against Bursa Malaysia's current official pages."
    )


def create_source_pdfs() -> None:
    create_product_overview_pdf()
    create_pdf(INPUT / "Learn" / "Bursa_Market_Primer.pdf", "Bursa Market Primer", "Plain-language starter material for new joiners", [
        ("1. What an exchange does", ["An exchange brings together investors and issuers within an organised market. It provides listing, trading, clearing and market-information services under applicable rules.", "BursaIQ should help a new joiner understand a term while keeping the original source one click away."]),
        ("2. Average Daily Value (ADV)", ["ADV is total traded value divided by the number of trading days in the measurement window. Always name the window because a 30-day ADV and a monthly ADV may differ.", "Worked synthetic example: RM68.4 billion traded over 20 trading days gives RM3.42 billion ADV."]),
        ("3. Market capitalisation", ["Market capitalisation is the market value of listed securities. At company level it is commonly price multiplied by issued shares. At market level the eligible company values are aggregated."]),
        ("4. Trading velocity", ["Trading velocity relates annualised trading value to market capitalisation. It indicates activity relative to market size; it does not by itself prove liquidity quality."]),
        ("5. Index attribution", ["Index attribution estimates how constituents or sectors added to or detracted from the index move. Results depend on the index method, constituent weights and reporting cut-off."]),
        ("6. Investor participation and net flow", ["Participation share shows how much traded value came from an investor group. Net flow is purchase value less sale value. Aggregated figures should avoid disclosing investor identity."])
    ])
    create_pdf(INPUT / "Learn" / "New_Joiner_Conduct_Guide.pdf", "New Joiner Conduct Guide", "Responsible information use in the BursaIQ demonstration", [
        ("1. Protect information", ["Use only approved material in BursaIQ. Do not paste confidential, personal or market-sensitive data into a demonstration environment."]),
        ("2. Check the source", ["Read the cited passage, date, owner and assumptions before reusing an answer. BursaIQ supports judgement; it does not replace the accountable subject-matter owner."]),
        ("3. Escalate uncertainty", ["Send a result to Verification Centre when it may inform a management decision, external statement or people decision."]),
        ("4. Communicate honestly", ["Label synthetic and illustrative figures. Do not present prototype access controls as production security."])
    ])
    create_pdf(INPUT / "HR" / "Hiring_Procedure_Demo.pdf", "Talent Acquisition Procedure — Demo", "Fictional procedure used to demonstrate the HR workspace", [
        ("1. Requisition approval", ["Owner: Hiring manager and Finance. Target: 2 working days. Confirm the business need, approved headcount and role profile."]),
        ("2. Sourcing and screening", ["Owner: Talent Acquisition. Target: 8 working days. Apply the approved selection criteria consistently and document the shortlist."]),
        ("3. Panel assessment", ["Owner: Hiring panel. Target: 5 working days. Use the agreed scorecard and record panel feedback. The next stage is pre-employment checks."]),
        ("4. Pre-employment checks", ["Owner: Talent Acquisition. Target: 4 working days. Candidate action time may pause the service clock."]),
        ("5. Offer approval and issue", ["Owner: HR approver. Target: 3 working days. Confirm approval before issuing the offer. Escalate overdue action to the named stage owner."])
    ])


def main() -> None:
    create_market_workbook()
    create_hr_workbook()
    create_source_pdfs()
    print(f"Synthetic demo pack created under {INPUT}")


if __name__ == "__main__":
    main()
